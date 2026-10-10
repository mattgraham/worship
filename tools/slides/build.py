# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=1.0"]
# ///
"""Render each set list to a lyrics-only projection deck (.pptx).

    uv run tools/slides/build.py                    # every sets/*.txt -> dist/<name>.pptx
    uv run tools/slides/build.py sets/MEG26.txt     # just that one
    uv run tools/slides/build.py --song "10000 Reasons"   # one song, for a quick look

The .onsong charts are the single source of truth for the words on the
screen as well as on the music stand, so this reads the same files, with the
same set-list lines and the same title matching as tools/setpdf/build.py (it
imports that module rather than copying it). Chords ([G]), the key/tempo
header and a "| key" suffix on a set line are all irrelevant on a projector
and are dropped.

Slide breaks, in order of precedence:

  1. A "# ---" line inside a section is an invisible marker: the slide ends
     there. OnSong does not show lines starting with "#" and ChordPro drops
     them, so the marker costs nothing in the apps or in the printed PDF.
  2. Without markers a section is one slide.
  3. Whatever is still longer than MAX_LINES lyric lines is cut into balanced
     chunks (9 lines become 3+3+3, not 4+4+1), because a lone orphan line on a
     slide reads as a mistake from the back of the room.

A section that is only a repeat cue ("Chorus" / "(Repeat Bridge)") shows the
slides of the section it points at, if that can be found earlier in the song;
otherwise there is nothing to project and it is skipped, not guessed at.

Output is 16:9, white text centred on black, with the section label small and
grey in the top-left corner and one title slide per song.
"""

from __future__ import annotations

import argparse
import importlib.util
import math
import os
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent

# Imported by path, under another name: this file is itself called "build",
# and as the script it sits first on sys.path, so a plain `import build`
# would hand back this very module instead of setpdf's.
_spec = importlib.util.spec_from_file_location("setpdf_build", HERE.parent / "setpdf" / "build.py")
setpdf = importlib.util.module_from_spec(_spec)
sys.modules["setpdf_build"] = setpdf
_spec.loader.exec_module(setpdf)

MAX_LINES = 4
GREY = (0x9A, 0x9A, 0x9A)
WHITE = (0xFF, 0xFF, 0xFF)


# ---------------------------------------------------------------------------
# Splitting -- pure functions, no pptx needed (see test_build.py)

BREAK = None  # sentinel in a list of lyric lines: the "# ---" marker


def lyric_of(raw: str) -> str | None:
    """A chart line as words only, or None if nothing is left to sing.

    Lines that are only chords and bar lines ("| [Ab] / / / |", "[G] [C]")
    vanish, since once the brackets are gone what remains is slashes and
    digits, which setpdf already knows how to recognise.
    """
    parts = setpdf.BRACKET.split(raw)
    lyric = "".join(parts[0::2])
    if len(parts) > 1 and setpdf.CHORD_LINE_REST.match(lyric):
        return None
    # "even--ing", "fore----vermore": hand-typed syllable stretches that mean
    # "sing it long" on a chart (seen in 10000 Reasons) and are just typos on
    # a screen. Single hyphens ("a-bundant-ly") are left alone: they are
    # indistinguishable from real compounds.
    lyric = re.sub(r"(?<=\w)-{2,}(?=\w)", "", lyric)
    lyric = re.sub(r"\s+", " ", lyric).strip()
    return lyric or None


def balanced_chunks(lines: list[str], limit: int = MAX_LINES) -> list[list[str]]:
    """Cut into the fewest chunks of at most `limit` lines, sizes as equal
    as possible and the longer ones first."""
    if not lines:
        return []
    k = math.ceil(len(lines) / limit)
    base, extra = divmod(len(lines), k)
    out, i = [], 0
    for n in range(k):
        size = base + (1 if n < extra else 0)
        out.append(lines[i:i + size])
        i += size
    return out


def split_section(lines: list[str | None], limit: int = MAX_LINES) -> list[list[str]]:
    """Lyric lines (with BREAK markers) of one section -> slides' worth of text."""
    groups: list[list[str]] = [[]]
    for ln in lines:
        if ln is BREAK:
            groups.append([])
        else:
            groups[-1].append(ln)
    return [c for g in groups for c in balanced_chunks(g, limit)]


def section_lines(section) -> list[str | None]:
    out: list[str | None] = []
    for ln in section.lines:
        if ln["kind"] == "break":
            out.append(BREAK)
        elif ln["kind"] == "raw":
            text = lyric_of(ln["text"])
            if text:
                out.append(text)
        # notes ({comment}) and tab blocks are not for the congregation
    return out


def find_target(sections, idx: int):
    """The earlier section a repeat cue at `idx` points back to, or None.
    Exact label first ("Chorus" -> "Chorus"), then a numbered one
    ("Chorus" -> "Chorus 2"); the nearest one wins."""
    want = setpdf.norm(sections[idx].repeat or "")
    if not want:
        return None
    earlier = [s for s in sections[:idx] if s.label and s.lines]
    for test in (lambda lab: lab == want, lambda lab: lab.startswith(want + " ")):
        for s in reversed(earlier):
            if test(setpdf.norm(s.label)):
                return s
    return None


def song_slides(sections) -> list[tuple[str | None, list[str]]]:
    """Sections -> [(label, lines)], one per slide."""
    out = []
    for i, s in enumerate(sections):
        src = s
        if s.repeat and not s.lines:
            src = find_target(sections, i)
            if src is None:
                continue
        for chunk in split_section(section_lines(src)):
            out.append((src.label, chunk))
    return out


# ---------------------------------------------------------------------------
# pptx

def font_size(lines: list[str]) -> int:
    """Largest size that keeps a full slide inside the frame: long lines
    wrap, so they cost more height."""
    longest = max((len(l) for l in lines), default=0)
    for size, cols in ((54, 34), (46, 40), (40, 48), (34, 58)):
        if longest <= cols:
            return size
    return 30


def build_deck(songs: list[dict], path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide():
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(0, 0, 0)
        return s

    def text(s, x, y, w, h, paras, size, rgb, align, anchor):
        tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            run = p.add_run()
            run.text = para
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(*rgb)
        return tf

    for song in songs:
        s = slide()
        text(s, 0.8, 2.4, 11.7, 1.5, [song["title"]], 60, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if song["artist"]:
            text(s, 0.8, 4.0, 11.7, 0.8, [song["artist"]], 28, GREY, PP_ALIGN.CENTER, MSO_ANCHOR.TOP)
        for label, lines in song["slides"]:
            s = slide()
            if label:
                text(s, 0.4, 0.25, 6, 0.5, [label], 20, GREY, PP_ALIGN.LEFT, MSO_ANCHOR.TOP)
            text(s, 0.8, 0.9, 11.7, 5.7, lines, font_size(lines), WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def song_data(path: Path) -> dict:
    chart = setpdf.parse_chart(path)
    meta = chart.meta
    return {
        "title": meta.get("title") or path.stem,
        "artist": meta.get("artist") or meta.get("author") or meta.get("writer") or None,
        "slides": song_slides(chart.sections),
    }


# ---------------------------------------------------------------------------
# Set lists -- the same lines and matching as setpdf


def build_set(name: str, lines: list[str], charts, out_dir: Path) -> Path | None:
    print(f"== {name}")
    songs, skipped = [], 0
    for line in lines:
        if re.match(r"^\s*#", line):
            continue
        title, _key, _bad = setpdf.parse_set_line(line)  # a key means nothing without chords
        q = setpdf.norm(title)
        if not q:
            continue
        hits = setpdf.resolve(q, charts)
        if len(hits) == 1:
            print(f"   ok       {line.strip()} -> {hits[0].name}")
            songs.append(song_data(hits[0]))
        elif not hits:
            print(f"   MISSING  {line.strip()}", file=sys.stderr)
            skipped += 1
        else:
            names = " ".join(f'"{p.name}"' for p in hits)
            print(f"   AMBIGUOUS {line.strip()} -> {names}", file=sys.stderr)
            skipped += 1

    if not songs:
        print("   no songs, no deck")
        return None
    out = out_dir / f"{name}.pptx"
    build_deck(songs, out)
    n = sum(len(s["slides"]) for s in songs)
    print(f"   {len(songs)} songs, {n} lyric slides -> {out}" + (f", {skipped} skipped" if skipped else ""))
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("sets", nargs="*", type=Path, help="set files (default: sets/*.txt)")
    ap.add_argument("--song", action="append", help='build one set line, e.g. "Oceans"')
    ap.add_argument("--out", type=Path, default=Path(os.environ.get("OUT_DIR", REPO / "dist")))
    args = ap.parse_args()

    charts = setpdf.index_charts()
    if not charts:
        print("no .onsong charts found", file=sys.stderr)
        return 1
    if args.song:
        build_set("song", args.song, charts, args.out)
        return 0
    for s in args.sets or sorted((REPO / "sets").glob("*.txt")):
        build_set(s.stem, s.read_text(encoding="utf-8").splitlines(), charts, args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
